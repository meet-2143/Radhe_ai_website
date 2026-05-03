from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Colors from index.css
COLOR_BG = RGBColor(5, 5, 10)       
COLOR_PRIMARY = RGBColor(118, 51, 255)  
COLOR_SECONDARY = RGBColor(0, 214, 255) 
COLOR_FG = RGBColor(250, 250, 250)     

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG

def add_accent_shapes(slide):
    shapes = slide.shapes
    
    # Top aesthetic stripe
    shape_top = shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(10), Inches(0.08))
    shape_top.fill.solid()
    shape_top.fill.fore_color.rgb = COLOR_PRIMARY
    shape_top.line.fill.background()

    # Bottom accent line
    shape_bot = shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.42), Inches(10), Inches(0.08))
    shape_bot.fill.solid()
    shape_bot.fill.fore_color.rgb = COLOR_SECONDARY
    shape_bot.line.fill.background()

    # Decorative side shape (subtle)
    shape_side = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.8), Inches(1), Inches(0.2), Inches(5.5))
    shape_side.fill.solid()
    shape_side.fill.fore_color.rgb = COLOR_PRIMARY
    shape_side.line.fill.background()

def add_slide_v4(prs, title_text, content_text=None, is_title_slide=False):
    if is_title_slide:
        slide_layout = prs.slide_layouts[0]
    else:
        slide_layout = prs.slide_layouts[1]
    
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_accent_shapes(slide)
    
    # Title Styling
    if slide.shapes.title:
        title = slide.shapes.title
        title.text = title_text
        title_para = title.text_frame.paragraphs[0]
        title_para.font.name = 'Arial' 
        title_para.font.size = Pt(44) if is_title_slide else Pt(32)
        title_para.font.bold = True
        title_para.font.color.rgb = COLOR_FG
        title_para.alignment = PP_ALIGN.CENTER if is_title_slide else PP_ALIGN.LEFT

    if is_title_slide:
        if len(slide.placeholders) > 1:
            subtitle = slide.placeholders[1]
            subtitle.text = content_text if content_text else ""
            sub_para = subtitle.text_frame.paragraphs[0]
            sub_para.font.name = 'Arial'
            sub_para.font.size = Pt(22)
            sub_para.font.color.rgb = COLOR_SECONDARY
            sub_para.alignment = PP_ALIGN.CENTER
    elif content_text:
        if len(slide.placeholders) > 1:
            content = slide.placeholders[1]
            content.text = content_text
            for para in content.text_frame.paragraphs:
                para.font.name = 'Arial'
                para.font.size = Pt(16)
                para.font.color.rgb = COLOR_FG
                para.space_after = Pt(10)

    return slide

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Title Slide
    add_slide_v4(prs, "RADHE AI", "ARCHITECTING THE FUTURE OF INTELLIGENCE\nCompany Presentation", is_title_slide=True)

    # Content Slides
    sections = [
        ("OUR STORY & MISSION", 
         "Founded to democratize high-end intelligence, Radhe AI provides global IT solutions.\n\n"
         "MISSION: To architect secure, intelligent, and scalable digital ecosystems.\n\n"
         "VISION: Harmony between AI and human ingenuity for global progress."),
        
        ("OUR LEADERSHIP", 
         "PARTH MAVANI - FOUNDER & CEO\nVisionary architect of complex AI systems.\n\n"
         "MEET MAVANI - CO-FOUNDER & CTO\nTechnical lead in Large-scale Infra and Gen AI."),
        
        ("CORE SERVICES (I)", 
         "• Answer Engine Optimization (AEO) for LLMs\n"
         "• Search Engine Optimization (SEO)\n"
         "• Digital Marketing & Performance ROI\n"
         "• Content Marketing & Authority\n"
         "• Advanced Analytics & GA4/GTM Reporting"),
        
        ("CORE SERVICES (II)", 
         "• Custom AI & Machine Learning Models\n"
         "• Zero-Trust Cloud Infrastructure (AWS/Azure/GCP)\n"
         "• Cyber Security & Risk Management\n"
         "• Premium UI/UX Digital Product Design"),
        
        ("SUCCESS STORIES (I)", 
         "• Fintech Scale: 10M+ Users, 99.999% Uptime\n"
         "• NeuralTrade V3: AI Trading Engine (<2ms Latency)\n"
         "• LogiFlow AI: Supply Chain Optimization ($4.5M Saved)\n"
         "• MedSafe Hub: Zero-Knowledge Medical Privacy"),
        
        ("SUCCESS STORIES (II)", 
         "• SecureFin: Cyber Overhaul for Tier 1 Banking\n"
         "• StreamlineCRM: 98% Reduction in Lead Response Time\n"
         "• NexusPOS: Event-driven Inventory for 500+ Sites\n"
         "• RankWave AI: 100% Traffic Recovery Via NLP"),
        
        ("MILESTONES", 
         "2021: Founded in a boutique tech hub.\n"
         "2022: Scaled to 50+ enterprise clients globally.\n"
         "2023: Awarded 'Top AI Innovator' by TechInsights.\n"
         "2024: Launched First Zero-Trust AI Infrastructure."),
        
        ("PARTNERS & TRUST", 
         "• AWS Partner Network // NVIDIA Inception\n"
         "• Google Cloud Partner // Microsoft Silver\n"
         "• ISO 27001 & SOC2 Certified Compliance"),
        
        ("CONTACT US", 
         "READY TO LEAD THE AI REVOLUTION?\n\n"
         "Visit: www.radhe-ai.com\n"
         "Email: contact@radhe-ai.com")
    ]

    for title, content in sections:
        add_slide_v4(prs, title, content)

    # Save
    prs.save('Radhe_AI_Presentation.pptx')
    print("Enhanced presentation successfully updated.")

if __name__ == "__main__":
    create_presentation()
